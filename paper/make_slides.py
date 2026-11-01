"""Build the CA3 video-presentation deck from the project's own figures and renders.

    python paper/make_slides.py          -> D:/dev/CA3-deliverables/Reactionless-Compliance-CA3-slides.pptx

Every slide names its presenter (P1-P5) so each member presents their own part (CA3 rule:
every member appears in the video). Speaker notes hold a ~40-60 s script per slide.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG, PFIG, V1 = ROOT / "figures", ROOT / "figures" / "paper", ROOT / "v1" / "paper" / "figs"
OUT = Path(r"D:\dev\CA3-deliverables\Reactionless-Compliance-CA3-slides.pptx")

NAVY, INK, PANEL = RGBColor(0x0B, 0x14, 0x26), RGBColor(0x13, 0x22, 0x3D), RGBColor(0x1B, 0x2D, 0x4F)
AMBER, LIGHT, MUTED = RGBColor(0xF2, 0xB7, 0x05), RGBColor(0xE6, 0xEC, 0xF5), RGBColor(0x9F, 0xB0, 0xC9)
RED, GREEN, WHITE = RGBColor(0xE5, 0x48, 0x4D), RGBColor(0x4C, 0xC3, 0x8A), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"
ROLE = {"P1": "P1 · Sim & integration", "P2": "P2 · Dynamics", "P3": "P3 · Control",
        "P4": "P4 · Safety", "P5": "P5 · Experiment", "ALL": "Team"}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def bg(slide, color=NAVY):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, s, size=18, color=LIGHT, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line, tuple):
            line, kw = line
        else:
            kw = {}
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size = font, Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", bold)
        r.font.color.rgb = kw.get("color", color)
        p.alignment = align
        p.space_after = Pt(kw.get("after", 6))
    return tb


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def header(slide, title, who, sub=None):
    rect(slide, Inches(0.5), Inches(0.45), Inches(0.09), Inches(0.75), AMBER)
    text(slide, Inches(0.75), Inches(0.38), Inches(9.7), Inches(0.8), title, size=27, bold=True, color=WHITE)
    if sub:
        text(slide, Inches(0.75), Inches(1.02), Inches(10), Inches(0.5), sub, size=15, color=MUTED)
    badge = rect(slide, Inches(10.6), Inches(0.45), Inches(2.25), Inches(0.5), PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = badge.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = ROLE[who]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(13), True, AMBER


def pic(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def pic_fit(slide, path, x, y, w, h):
    """Largest picture with preserved aspect ratio inside the box, centred."""
    from PIL import Image
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = int(iw * s), int(ih * s)
    return slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def card(slide, x, y, w, h, title, body, accent=AMBER, size=15):
    rect(slide, x, y, w, h, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y + Inches(0.15), Inches(0.06), h - Inches(0.3), accent)
    text(slide, x + Inches(0.2), y + Inches(0.1), w - Inches(0.3), Inches(0.45), title, size=size + 2, bold=True, color=WHITE)
    text(slide, x + Inches(0.2), y + Inches(0.55), w - Inches(0.3), h - Inches(0.6), body, size=size, color=LIGHT)


def stat(slide, x, y, w, big, small, color=AMBER):
    text(slide, x, y, w, Inches(0.9), big, size=40, bold=True, color=color)
    text(slide, x, y + Inches(0.85), w, Inches(0.8), small, size=14, color=MUTED)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def footer(slide, n):
    text(slide, Inches(0.5), Inches(7.0), Inches(9), Inches(0.35),
         "Reactionless Compliance · MA6223 CA3 · NTU", size=10, color=MUTED)
    text(slide, Inches(12.2), Inches(7.0), Inches(0.7), Inches(0.35), str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


slides = []


def new():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    slides.append(s)
    return s


# 1 -- title -----------------------------------------------------------------------------
s = new()
pic_fit(s, FIG / "render_lineup.png", Inches(0), Inches(3.3), W, Inches(4.0))
rect(s, Inches(0.5), Inches(0.9), Inches(0.12), Inches(1.9), AMBER)
text(s, Inches(0.8), Inches(0.7), Inches(12), Inches(1.4), "Reactionless Compliance", size=50, bold=True, color=WHITE)
text(s, Inches(0.8), Inches(1.65), Inches(12), Inches(1.2),
     "A human-in-the-loop benchmark of redundancy resolution and contact-safety limits "
     "for teleoperated free-floating space manipulators", size=20, color=LIGHT)
text(s, Inches(0.8), Inches(2.6), Inches(12), Inches(0.5),
     "MA6223 CA3 · MSc Robotics and Intelligent Systems · Nanyang Technological University", size=14, color=MUTED)
notes(s, "[All, 15 s] Hello, we are group <number>. Each of us will present our own part: P1 the simulation and "
         "integration, P2 the dynamics, P3 the controllers, P4 contact safety and P5 the experiment and results. "
         "The picture shows our three simulated servicers, from a 20 kg CubeSat-class base to a 1500 kg one, "
         "all carrying the same 7-joint Kinova Gen3 arm.")

# 2 -- the problem ---------------------------------------------------------------------
s = new()
header(s, "The problem: an arm that turns its own spacecraft", "P5")
card(s, Inches(0.6), Inches(1.6), Inches(3.9), Inches(2.4), "Swivel-chair demo",
     "Sit on an office chair, feet up, swing your arm in a circle: you rotate. A servicer arm does the "
     "same to its satellite, and the antenna loses the ground station.")
card(s, Inches(0.6), Inches(4.2), Inches(3.9), Inches(2.4), "Elbow demo",
     "Trace a circle with a fingertip and return to the same dot: the elbow has moved. Redundant arms "
     "drift in posture even when the hand comes home.", accent=MUTED)
rect(s, Inches(4.9), Inches(1.6), Inches(7.9), Inches(5.0), PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(5.2), Inches(1.75), Inches(7.4), Inches(0.5), "Why it cannot be solved, only traded", size=20, bold=True, color=WHITE)
text(s, Inches(5.2), Inches(2.45), Inches(7.4), Inches(2.0),
     [("7 joints", {"size": 34, "bold": True, "color": WHITE}),
      ("− 6  hand position + orientation (6-D task)", {"size": 22, "color": LIGHT}),
      ("− 3  keep the base attitude still", {"size": 22, "color": LIGHT}),
      ("= −2   over-constrained", {"size": 28, "bold": True, "color": RED})])
text(s, Inches(5.2), Inches(5.25), Inches(7.4), Inches(1.3),
     "Every controller must trade task accuracy, base drift and posture drift. Northrop Grumman's MRV "
     "(launched 2026) has human-operated arms: this is the operator's problem today.", size=15, color=MUTED)
notes(s, "[P5, 45 s] Two quick demos explain the problem. On a swivel chair, swinging your arm turns your "
         "body; in space, an arm turns its satellite and the antenna drifts off the ground station. And when "
         "you trace a circle and come back, your elbow is not where it started: posture drift. A 7-joint arm "
         "doing a full 6-D task has one spare joint, but holding the base still needs three. So exact "
         "reactionless motion is impossible, and every controller makes a trade-off. Our project measures it.")

# 3 -- what we built ---------------------------------------------------------------------
s = new()
header(s, "What we built", "P5", "One open benchmark, four contributions")
items = [("1  Benchmark", "3 servicers × 5 schemes × 3-D/6-D, scripted ORU task, fixed tuning rule, β-sweep", AMBER),
         ("2  Trade-off", "Task error vs base drift vs posture drift, and where ±5° antenna pointing breaks", AMBER),
         ("3  Contact safety", "ISO/TS 15066 body model where ISO 10218 excludes space; validated limiter", GREEN),
         ("4  Human-in-the-loop", "PS4 interface with haptics, counterbalanced protocol, analysis pipeline", MUTED)]
for i, (t, b, c) in enumerate(items):
    card(s, Inches(0.6 + (i % 2) * 6.2), Inches(1.7 + (i // 2) * 2.55), Inches(6.0), Inches(2.3), t, b, accent=c, size=16)
notes(s, "[P5, 30 s] We built an open benchmark with four parts: three servicer sizes and five control schemes in "
         "position-only and full-pose tasks; a characterisation of the trade-off; a contact-safety limit, "
         "because the industrial robot standard explicitly excludes space robots; and a teleoperation "
         "interface and study protocol so that real operators can compare the schemes.")

# 4 -- servicers and arm (P1) ---------------------------------------------------------------
s = new()
header(s, "Three servicers, one Kinova Gen3 arm", "P1", "MuJoCo, momentum-conserving, gravity off")
pic_fit(s, FIG / "render_medium.png", Inches(0.5), Inches(1.6), Inches(7.2), Inches(4.5))
rows = [("Servicer base", "21.7 / 151.7 / 1501.7 kg"), ("Arm / base mass ratio", "0.30 / 0.043 / 0.0043"),
        ("Arm", "Kinova Gen3 7-DoF, 902 mm, 6.5 kg"), ("Joints", "velocity servos, ±1 rad/s, ±5 Nm"),
        ("Integrator", "RK4, 0.25 ms; control 100 Hz"), ("Collisions", "arm–base and arm–arm on")]
for i, (k, v) in enumerate(rows):
    y = Inches(1.75 + i * 0.75)
    text(s, Inches(8.0), y, Inches(2.4), Inches(0.5), k, size=14, color=MUTED)
    text(s, Inches(10.2), y, Inches(2.8), Inches(0.6), v, size=15, bold=True, color=WHITE)
notes(s, "[P1, 45 s] I built the simulation. Three servicers share one arm and differ only in the base, so the "
         "arm-to-base mass ratio is the single variable: 0.30, 0.043 and 0.0043. The arm uses the real Kinova "
         "Gen3 geometry and masses from MuJoCo Menagerie. Joints are velocity servos with realistic torque "
         "saturation. We use a small RK4 time step because the momentum check must pass at one part in a "
         "million. The yellow cylinder is the antenna; it must stay within five degrees.")

# 5 -- teleop interface (P1/P3) --------------------------------------------------------------
s = new()
header(s, "Teleoperation: PS4 pad, haptics, head-up display", "P1")
pic_fit(s, FIG / "screenshot_teleop.png", Inches(0.5), Inches(1.5), Inches(7.6), Inches(5.3))
card(s, Inches(8.4), Inches(1.5), Inches(4.4), Inches(2.6), "Game-style layout",
     ["Left stick: move target (camera-relative)", "R2 / L2: up / down", "Right stick: orbit camera",
      "Hold R1: rotate · Hold L1: fine"], size=14)
card(s, Inches(8.4), Inches(4.3), Inches(4.4), Inches(2.5), "Haptics (DualShock 4)",
     ["Heavy motor: arm contact, scaled by force", "Light motor: last 15° before a joint limit",
      "Pulse: antenna leaves ±5° cone"], accent=RED, size=14)
notes(s, "[P1, 45 s] This is what the operator sees: the blue sphere is the commanded target, the orange ball is "
         "the next waypoint of the orbital-replacement-unit task, and the antenna lines show attitude drift. "
         "The PS4 controller uses a game-style layout, so classmates need little training. The rumble motors "
         "warn about contact with the spacecraft, joint limits and the antenna leaving its cone. A normal task "
         "never triggers them, so a buzz always means something is wrong. Everything is logged at 100 Hz.")

# 6 -- dynamics (P2) ---------------------------------------------------------------------------
s = new()
header(s, "Free-floating dynamics", "P2", "Generalized Jacobian and the reaction null space")
rect(s, Inches(0.6), Inches(1.7), Inches(7.2), Inches(4.0), PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(0.9), Inches(1.9), Inches(6.8), Inches(4.6),
     [("Momentum from rest", {"size": 16, "color": MUTED}),
      ("H_b v_b + H_bm q̇ = 0   ⇒   v_b = V q̇", {"size": 24, "bold": True, "color": WHITE, "after": 14}),
      ("Hand motion", {"size": 16, "color": MUTED}),
      ("ẋ = (J_m + J_b V) q̇ = J* q̇", {"size": 24, "bold": True, "color": WHITE, "after": 14}),
      ("Base does not rotate when", {"size": 16, "color": MUTED}),
      ("W q̇ = 0,   W = angular rows of V", {"size": 24, "bold": True, "color": AMBER})])
stat(s, Inches(8.3), Inches(1.8), Inches(4.5), "6 × 10⁻⁸", "worst momentum residual, 10 s of random motion, all 3 servicers (gate: 10⁻⁶)", GREEN)
stat(s, Inches(8.3), Inches(3.6), Inches(4.5), "rank 7 of 9", "6-D task + reactionless rows: infeasible (3-D: rank 6 of 6, exact)", AMBER)
stat(s, Inches(8.3), Inches(5.4), Inches(4.5), "10⁻⁷", "J* and W match simulated hand and base velocities", GREEN)
notes(s, "[P2, 50 s] With the base floating and starting at rest, momentum conservation gives the base velocity "
         "as a linear function of joint rates. Substituting into the hand kinematics gives the generalized "
         "Jacobian J-star. The angular rows W tell us how joints rotate the base; joint motions with W q-dot "
         "equal to zero are reactionless. Our momentum gate runs in CI on all three servicers: the residual "
         "stays at six times ten to the minus eight. And numerically, the 6-D task plus three reactionless rows "
         "has rank seven out of nine: the infeasibility we started from is real.")

# 7 -- schemes (P3) -----------------------------------------------------------------------------
s = new()
header(s, "Five redundancy-resolution schemes", "P3", "Same pose servo, same joint geofence, tuned with a rule fixed in advance")
sch = [("1  DLS", "q̇ = J*ᵀ(J*J*ᵀ + λ²I)⁻¹ ẋ", "λ = 0.01"),
       ("2  Gradient projection", "q̇ = J*⁺ẋ + (I − J*⁺J*) k_g ∇w", "k_g = 1"),
       ("3  Impedance", "null-space pull to rest posture", "K_p = 2"),
       ("4  RNS", "3-D: [J*; W] q̇ exact · 6-D: weighted LS", "k_a = 1"),
       ("5  Impedance + RNS", "posture pull inside the reaction null space", "K_p = 1")]
for i, (n, eq, g) in enumerate(sch):
    y = Inches(1.65 + i * 0.78)
    rect(s, Inches(0.6), y, Inches(12.2), Inches(0.68), PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(0.8), y + Inches(0.12), Inches(3.4), Inches(0.5), n, size=17, bold=True, color=AMBER if i > 2 else WHITE)
    text(s, Inches(4.2), y + Inches(0.12), Inches(6.4), Inches(0.5), eq, size=16, color=LIGHT)
    text(s, Inches(10.8), y + Inches(0.12), Inches(1.9), Inches(0.5), g, size=15, color=MUTED)
text(s, Inches(0.6), Inches(5.65), Inches(12.2), Inches(1.2),
     [("6-D:  min ‖J*q̇ − ẋ‖² + β ‖W q̇ + k_a φ‖² + γ ‖q̇ − q̇₀‖²", {"size": 20, "bold": True, "color": WHITE}),
      ("β trades task accuracy for base stillness · tuning: 40 runs, rule fixed before looking (no strawmen)", {"size": 14, "color": MUTED})])
notes(s, "[P3, 55 s] I implemented and tuned the five schemes. Damped least squares and gradient projection are "
         "fixed-base classics; impedance pulls the arm back to its rest posture in the null space. RNS solves the "
         "task and the reactionless rows together, and scheme five adds the posture pull inside the reaction "
         "null space. In 6-D the rows do not fit, so schemes four and five become a weighted least-squares "
         "problem, and beta sets the trade-off. To avoid strawmen, each scheme's own gain was swept and picked "
         "by a rule we fixed before looking at results. All five share the same target servo and joint geofence.")

# 8 -- trade-off (P3) ------------------------------------------------------------------------------
s = new()
header(s, "Result 1 — the 6-D trade-off", "P3", "β-sweep, scripted two-loop ORU task")
pic_fit(s, PFIG / "fig1_tradeoff.png", Inches(0.4), Inches(1.5), Inches(12.5), Inches(3.9))
stat(s, Inches(0.8), Inches(5.45), Inches(4), "1.3–2.8×", "less peak drift, schemes 4–5 (medium)")
stat(s, Inches(5.0), Inches(5.45), Inches(4), "8–22 mm", "the price: hand error vs 0.6 mm for DLS", RED)
stat(s, Inches(9.2), Inches(5.45), Inches(3.8), "β ≤ 3", "most reactionless setting that still finishes", GREEN)
notes(s, "[P3, 45 s] Each curve is the beta-sweep for schemes four and five; the single points are schemes one to "
         "three. On the medium servicer, raising beta moves the base from 1.7 to 0.3 degrees of drift but costs "
         "up to 43 millimetres of hand error, and above beta equals three the hand no longer meets the 2 cm "
         "tolerance. On the large servicer every scheme sits below 0.08 degrees, so the choice does not matter. "
         "A mission can read the right beta off this curve from its pointing budget.")

# 9 -- size + demo video (P3/P5) ---------------------------------------------------------------------
s = new()
header(s, "Result 2 — when does drift matter?", "P5", "Same scripted input on the 20 kg servicer: DLS vs Imp.+RNS")
mv = FIG / "demo_dls_vs_imprns.mp4"
if mv.exists():
    s.shapes.add_movie(str(mv), Inches(0.5), Inches(1.6), Inches(7.6), Inches(3.41),
                       poster_frame_image=str(FIG / "demo_frame.png"), mime_type="video/mp4")
else:
    pic_fit(s, FIG / "demo_frame.png", Inches(0.5), Inches(1.6), Inches(7.6), Inches(3.41))
pic_fit(s, PFIG / "fig2_size.png", Inches(8.3), Inches(1.6), Inches(4.6), Inches(2.2))
text(s, Inches(8.3), Inches(3.9), Inches(4.6), Inches(3.0),
     [("20 kg (ratio 0.30)", {"size": 17, "bold": True, "color": RED}),
      ("Schemes 1–3 finish but tip the antenna 23–25°; schemes 4–5 stay < 1.5° but cannot finish.", {"size": 14}),
      ("150 kg (0.043)", {"size": 17, "bold": True, "color": AMBER}),
      ("All finish within ±5°: the reactionless schemes earn their keep.", {"size": 14}),
      ("1500 kg (0.0043)", {"size": 17, "bold": True, "color": GREEN}),
      ("< 0.08° for every scheme: no need for reactionless control.", {"size": 14})])
text(s, Inches(0.5), Inches(5.15), Inches(7.6), Inches(1.6),
     "Left: DLS follows the loop but the antenna line turns red (past 5°). Right: Imp.+RNS keeps the "
     "base almost still but cannot reach the waypoints. Gate: 26/30 runs finish; medium and large 10/10.",
     size=14, color=MUTED)
notes(s, "[P5, 50 s] Play the clip. Both arms get exactly the same operator input on the 20 kg servicer. Damped "
         "least squares follows the loop, but the antenna line turns red: the base tips past five degrees and "
         "the radio link would be lost. Impedance plus RNS keeps the base almost still, but it cannot reach the "
         "waypoints. So for this arm a CubeSat-class servicer cannot do the task within the pointing limit with "
         "any scheme. On 150 kg every scheme works and the reactionless ones reduce drift; on 1500 kg drift "
         "is negligible.")

# 10 -- posture drift (P3) --------------------------------------------------------------------------------
s = new()
header(s, "Result 3 — the hand comes home, the arm does not", "P3")
pic_fit(s, PFIG / "fig3_joint_return.png", Inches(0.4), Inches(1.5), Inches(12.5), Inches(4.0))
card(s, Inches(0.6), Inches(5.6), Inches(6.0), Inches(1.25), "Without a posture term",
     "Plain RNS ends 2.92 rad from its start (3-D, medium) and its base 0.059° off: our one failed requirement (R4).",
     accent=RED, size=14)
card(s, Inches(6.8), Inches(5.6), Inches(6.0), Inches(1.25), "With posture impedance",
     "Imp.+RNS: 0.065 rad and 0.001°. Posture impedance makes reactionless control repeatable.",
     accent=GREEN, size=14)
notes(s, "[P3, 40 s] After two loops the hand is back, but the arm is not. Without a posture term, plain RNS lets "
         "the joints wander through the null space: almost three radians in 3-D, and the base ends 0.059 degrees "
         "off, which fails our 0.05-degree requirement. We kept that as a fail rather than move the threshold. "
         "Adding posture impedance inside the reaction null space, scheme five, brings it to 0.065 radians and "
         "0.001 degrees. For an operator, that repeatability matters.")

# 11 -- safety model (P4) ---------------------------------------------------------------------------------
s = new()
header(s, "Contact safety: which mass hits the astronaut?", "P4", "ISO/TS 15066 body model; ISO 10218-1:2025 excludes space robots")
pic_fit(s, PFIG / "fig5_vmax.png", Inches(0.4), Inches(1.55), Inches(7.6), Inches(5.3))
text(s, Inches(8.2), Inches(1.65), Inches(4.7), Inches(5.2),
     [("v_max = F / √(μ k)", {"size": 22, "bold": True, "color": WHITE}),
      ("All ISO body regions, free (recoil) and clamped (foot restraint) contact.", {"size": 14, "after": 12}),
      ("Rigid joints: the whole spacecraft hits", {"size": 16, "bold": True, "color": RED}),
      ("1473 kg on the large servicer → 13 mm/s limit (top row).", {"size": 14, "after": 12}),
      ("Real ±5 Nm servos give way", {"size": 16, "bold": True, "color": GREEN}),
      ("432 simulated impacts: across the arm the base mass barely matters → ~0.15–0.18 m/s (bottom row).",
       {"size": 14, "after": 12}),
      ("Joint compliance, not posture, sets the limit.", {"size": 16, "bold": True, "color": AMBER})])
notes(s, "[P4, 55 s] Industrial robot safety standards exclude space robots, so we borrowed the ISO/TS 15066 body "
         "model: the safe speed depends on the effective mass that hits each body region. With rigid joints the "
         "whole spacecraft stands behind the hand: 1473 kilograms on the large servicer, which would allow only "
         "13 millimetres per second, the top row. But we simulated 432 impacts into a body model: the five "
         "newton-metre servos give way, so across the arm the base mass barely matters. The real limit is about "
         "0.15 to 0.18 metres per second and almost flat over the workspace, the bottom row.")

# 12 -- limiter + verification (P4) --------------------------------------------------------------------------
s = new()
header(s, "A speed limiter we tested by crashing the arm", "P4")
pic_fit(s, PFIG / "figS1_safety.png", Inches(0.4), Inches(1.4), Inches(7.4), Inches(3.2))
stat(s, Inches(0.8), Inches(4.75), Inches(3.4), "397 / 397", "impacts at limiter speed below the ISO force limit (worst 0.90)", GREEN)
stat(s, Inches(4.4), Inches(4.75), Inches(3.4), "100 %", "of poses where a 5 kg payload exceeds the clamped limit without it", RED)
req = [("R1 antenna ±5°", "med ✓  small ✗"), ("R2 contact speed ≤ v_max", "✓"), ("R3 momentum < 10⁻⁶", "✓"),
       ("R4 3-D reactionless ≈ 0", "✗ plain RNS"), ("R5 joint limits", "✓"), ("R6 contact signalled", "✓"),
       ("R7 joint rates ≤ 1 rad/s", "✓"), ("R8 DOF argument", "✓"), ("R9 head / face", "procedural"),
       ("R10 input loss → zero", "✓")]
rect(s, Inches(8.1), Inches(1.4), Inches(4.75), Inches(5.4), PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(8.3), Inches(1.5), Inches(4.4), Inches(0.5), "Verification matrix (from data)", size=16, bold=True, color=WHITE)
for i, (k, v) in enumerate(req):
    y = Inches(2.05 + i * 0.465)
    text(s, Inches(8.3), y, Inches(2.9), Inches(0.45), k, size=13, color=LIGHT)
    text(s, Inches(11.1), y, Inches(1.7), Inches(0.45), v, size=13, bold=True,
         color=GREEN if v == "✓" else (RED if "✗" in v else AMBER))
notes(s, "[P4, 55 s] At run time the limiter scales the joint rates so the hand stays below 85 percent of the "
         "pose-dependent limit. We validated it directly: every simulated impact was re-run at exactly the "
         "limiter's speed, and all 397 contacts stayed below the ISO force limit; the worst reached 90 percent, "
         "a clamped hand pushed along the arm. A Monte Carlo sweep shows the limiter is not optional: with a 5 "
         "kilogram payload the interface's own top speed exceeds the clamped limit in every pose. On the right is "
         "our verification matrix, computed from the data: two honest fails, the small servicer and plain RNS.")

# 13 -- tilt-limit map (P2) -----------------------------------------------------------------------------------
s = new()
header(s, "Operator aid: see the tilt limit before you command", "P2", "Earlier v1 arm model (12.5 kg, 300 kg servicer)")
pic_fit(s, V1 / "fig_opview.png", Inches(0.5), Inches(1.6), Inches(7.4), Inches(5.0))
stat(s, Inches(8.3), Inches(1.7), Inches(4.5), "271 / 300", "targets near the edge classified right by the 98-probe 2° cage")
stat(s, Inches(8.3), Inches(3.4), Inches(4.5), "2.02°", "worst base tilt with the runtime fence, 46 refused commands", GREEN)
stat(s, Inches(8.3), Inches(5.1), Inches(4.5), "150–600 kg", "band of servicer sizes where the 2° limit binds", AMBER)
notes(s, "[P2, 45 s] In parallel I asked whether the operator can see the limit before commanding. Probing 98 "
         "directions from rest gives three nested cages for 0.5, 1 and 2 degrees of base tilt, drawn around the "
         "hand. On 300 targets near the edge the 2-degree cage was right on 271. A runtime fence then held the "
         "base at most 2.02 degrees, so a wrong command costs time, not attitude. And the limit only binds for "
         "servicers of roughly 150 to 600 kilograms, which matches the benchmark's size result.")

# 14 -- human study (P5) -----------------------------------------------------------------------------------
s = new()
header(s, "Human-in-the-loop study: ready to run", "P5")
steps = [("Design", "Within-subject: 5 schemes × 3-D/6-D on the medium servicer, 2 loops per trial"),
         ("Order", "Mode blocks alternate by participant; Williams Latin square for schemes; scheme hidden"),
         ("Measures", "Task error, base drift, joint drift, time · predictability (1–7) · raw NASA-TLX"),
         ("Statistics", "Friedman + Kendall's W; Wilcoxon signed-rank, Holm-corrected; effect size r"),
         ("Status", "Session runner, questionnaire, statistics and figures tested end to end; sessions pending")]
for i, (k, v) in enumerate(steps):
    y = Inches(1.6 + i * 1.0)
    c = rect(s, Inches(0.6), y, Inches(0.75), Inches(0.75), AMBER if i < 4 else MUTED, MSO_SHAPE.OVAL)
    c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = c.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb, r.font.name = str(i + 1), Pt(20), True, NAVY, FONT
    text(s, Inches(1.6), y + Inches(0.02), Inches(2.4), Inches(0.6), k, size=19, bold=True, color=WHITE)
    text(s, Inches(4.0), y + Inches(0.05), Inches(8.8), Inches(0.7), v, size=16, color=LIGHT)
notes(s, "[P5, 40 s] The study is designed and the tools are ready. Each participant flies all five schemes in "
         "both task modes on the medium servicer, in a counterbalanced Williams Latin-square order, without "
         "knowing which scheme is active. We log performance every tick and collect predictability ratings and "
         "NASA-TLX after each trial, analysed with Friedman and Holm-corrected Wilcoxon tests. Everything has "
         "been tested end to end; when this video was recorded the participant sessions had not yet been run.")

# 15 -- conclusion (ALL) -------------------------------------------------------------------------------------
s = new()
header(s, "Take-aways", "ALL")
take = [("Mass ratio decides the scheme", "0.30: none meets ±5° · 0.043: reactionless schemes help · 0.0043: any scheme"),
        ("Posture impedance makes reactionless control repeatable", "plain RNS drifts 2.9 rad; Imp.+RNS 0.07 rad"),
        ("Joint compliance shields the astronaut", "limiter validated on 397 simulated impacts, worst 90 % of ISO limit"),
        ("Limitations", "simulation only, no participants yet, one arm, ±5 Nm drives (real Gen3 is stronger)")]
for i, (t, b) in enumerate(take):
    card(s, Inches(0.6), Inches(1.45 + i * 1.22), Inches(12.2), Inches(1.08), t, b,
         accent=[AMBER, GREEN, GREEN, MUTED][i], size=16)
text(s, Inches(0.6), Inches(6.4), Inches(12), Inches(0.4),
     "Code, data and every number: github.com/Gonzalez-Annan/reactionless-compliance", size=13, color=AMBER)
notes(s, "[All, 30 s] Three take-aways. The arm-to-base mass ratio decides which scheme is right. Posture impedance "
         "is what makes reactionless control repeatable. And back-drivable joints shield the astronaut from the "
         "spacecraft's mass, with a limiter we validated on simulated impacts. Everything is simulated so far, "
         "and the human study comes next. Thank you.")

for i, sl in enumerate(slides[1:], 2):
    footer(sl, i)
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT, len(slides), "slides")
